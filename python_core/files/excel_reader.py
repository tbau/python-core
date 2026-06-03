"""Excel reader."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from python_core.exceptions import ConfigurationError, ValidationError
from python_core.files.excel_profile import ExcelProfile


class ExcelReader:
    """Read rows from an Excel workbook according to an ExcelProfile."""

    def __init__(self, profile: ExcelProfile) -> None:
        self.profile = profile

    def parse(self, path: Path) -> list[dict[str, Any]]:
        """Read workbook rows as dictionaries."""
        openpyxl = self._openpyxl()
        workbook = openpyxl.load_workbook(path, data_only=True)
        try:
            if self.profile.sheet_name not in workbook.sheetnames:
                raise ValidationError(f"missing sheet: {self.profile.sheet_name}")

            sheet = workbook[self.profile.sheet_name]
            headers = self._headers(sheet)
            missing = self.profile.required_columns.difference(headers)
            if missing:
                raise ValidationError(f"missing required columns: {sorted(missing)}")

            rows: list[dict[str, Any]] = []
            for values in sheet.iter_rows(min_row=self.profile.header_row + 1, values_only=True):
                row = dict(zip(headers, values, strict=False))
                if self.profile.row_filter and not self.profile.row_filter(row):
                    continue
                rows.append(row)
            return rows
        finally:
            workbook.close()

    def _headers(self, sheet: Any) -> list[str]:
        try:
            raw_headers = next(
                sheet.iter_rows(min_row=self.profile.header_row, max_row=self.profile.header_row)
            )
        except StopIteration as exc:
            raise ValidationError(f"missing header row: {self.profile.header_row}") from exc
        return [self.profile.normalize_header(str(cell.value or "")) for cell in raw_headers]

    def _openpyxl(self) -> Any:
        try:
            import openpyxl
        except ImportError as exc:
            raise ConfigurationError("Install python-core[excel] to read Excel files") from exc
        return openpyxl
