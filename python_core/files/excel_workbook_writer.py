"""Customizable Excel workbook writer."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from python_core.exceptions import ConfigurationError
from python_core.files.spreadsheet_column import SpreadsheetColumn
from python_core.files.spreadsheet_sheet import SpreadsheetSheet
from python_core.files.spreadsheet_style import SpreadsheetStyle
from python_core.files.spreadsheet_workbook import SpreadsheetWorkbook


class ExcelWorkbookWriter:
    """Writes a SpreadsheetWorkbook using openpyxl."""

    def __init__(self, *, overwrite: bool = True, create_parent_dirs: bool = True) -> None:
        self.overwrite = overwrite
        self.create_parent_dirs = create_parent_dirs

    def write(self, workbook: SpreadsheetWorkbook, path: Path) -> Path:
        """Write a workbook and return the output path."""
        if path.exists() and not self.overwrite:
            raise FileExistsError(path)
        if self.create_parent_dirs:
            path.parent.mkdir(parents=True, exist_ok=True)

        openpyxl = self._openpyxl()
        excel_workbook = openpyxl.Workbook()
        self._write_sheets(openpyxl, excel_workbook, workbook)
        if workbook.active_sheet and workbook.active_sheet in excel_workbook.sheetnames:
            excel_workbook.active = excel_workbook.sheetnames.index(workbook.active_sheet)
        excel_workbook.save(path)
        return path

    def _write_sheets(self, openpyxl: Any, excel_workbook: Any, workbook: SpreadsheetWorkbook) -> None:
        default_sheet = excel_workbook.active
        for index, sheet_model in enumerate(workbook.sheets):
            sheet = default_sheet if index == 0 else excel_workbook.create_sheet()
            sheet.title = sheet_model.name
            self._write_sheet(openpyxl, sheet, sheet_model, workbook.styles or {})

    def _write_sheet(
        self,
        openpyxl: Any,
        sheet: Any,
        sheet_model: SpreadsheetSheet,
        styles: dict[str, SpreadsheetStyle],
    ) -> None:
        rows = sheet_model.materialized_rows()
        columns = sheet_model.resolved_columns(rows)
        sheet.append([column.label for column in columns])
        self._style_header(openpyxl, sheet, sheet_model.header_style)

        for row in rows:
            sheet.append([column.value_for(row) for column in columns])

        self._apply_sheet_options(sheet, sheet_model, columns)
        self._apply_column_styles(openpyxl, sheet, columns, styles)

    def _style_header(self, openpyxl: Any, sheet: Any, style: SpreadsheetStyle) -> None:
        for cell in sheet[1]:
            self._apply_style(openpyxl, cell, style)

    def _apply_sheet_options(
        self,
        sheet: Any,
        sheet_model: SpreadsheetSheet,
        columns: list[SpreadsheetColumn],
    ) -> None:
        if sheet_model.freeze_panes:
            sheet.freeze_panes = sheet_model.freeze_panes
        if sheet_model.auto_filter and columns:
            sheet.auto_filter.ref = sheet.dimensions
        if sheet_model.tab_color:
            sheet.sheet_properties.tabColor = sheet_model.tab_color
        for index, column in enumerate(columns, start=1):
            if column.width:
                letter = sheet.cell(row=1, column=index).column_letter
                sheet.column_dimensions[letter].width = column.width

    def _apply_column_styles(
        self,
        openpyxl: Any,
        sheet: Any,
        columns: list[SpreadsheetColumn],
        styles: dict[str, SpreadsheetStyle],
    ) -> None:
        for index, column in enumerate(columns, start=1):
            if not column.style_key or column.style_key not in styles:
                continue
            for cell in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for item in cell:
                    self._apply_style(openpyxl, item, styles[column.style_key])

    def _apply_style(self, openpyxl: Any, cell: Any, style: SpreadsheetStyle) -> None:
        if style.bold or style.font_color:
            cell.font = openpyxl.styles.Font(bold=style.bold, color=style.font_color)
        if style.fill_color:
            cell.fill = openpyxl.styles.PatternFill("solid", fgColor=style.fill_color)
        if style.number_format:
            cell.number_format = style.number_format

    def _openpyxl(self) -> Any:
        try:
            import openpyxl
        except ImportError as exc:
            raise ConfigurationError("Install python-core[excel] to write Excel files") from exc
        return openpyxl
